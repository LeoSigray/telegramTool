"""Оценка качества по вашей разметке в Excel.

Поставьте «да» или «нет» в столбце «Разметка (да/нет)» на листе «Люди»
(хотя бы у 50–100 строк подряд сверху) и запустите:

    python -m parsing.smart.evaluate data/smart_leads/<файл>.xlsx

Выводит precision@k, точность по уровням и источникам и средние факторы
у «да» и «нет», чтобы видеть, какой фактор врёт.
"""
from __future__ import annotations

import sys
from collections import defaultdict

YES = {"да", "yes", "y", "1", "+", "д"}
NO = {"нет", "no", "n", "0", "-", "н"}
FACTORS = ("PQI", "I намерение", "R тема", "F свежесть", "T ветка", "E активность", "A ЛПР",
           "CQI источника")


def load_labels(path: str) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Люди"]
    rows = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(rows)]
    out = []
    for vals in rows:
        rec = dict(zip(header, vals))
        mark = str(rec.get("Разметка (да/нет)") or "").strip().lower()
        rec["_label"] = True if mark in YES else (False if mark in NO else None)
        out.append(rec)
    return out


def report(path: str) -> str:
    rows = load_labels(path)
    labeled = [r for r in rows if r["_label"] is not None]
    lines = [f"Файл: {path}", f"Строк: {len(rows)}, размечено: {len(labeled)}"]
    if not labeled:
        lines.append("Нет разметки: поставьте «да»/«нет» в столбце «Разметка (да/нет)».")
        return "\n".join(lines)

    lines.append("\nPrecision@k (доля «да» среди первых k размеченных по порядку PQI):")
    for k in (10, 25, 50, 100, 200):
        top = labeled[:k]
        if len(top) < k and k > 10:
            break
        lines.append(f"  @{k:<4} {sum(r['_label'] for r in top) / len(top):.0%}  ({len(top)} строк)")

    for key, title in (("Уровень", "по уровням"), ("Источник", "по источникам")):
        groups: dict = defaultdict(list)
        for r in labeled:
            groups[r.get(key)].append(r["_label"])
        lines.append(f"\nТочность {title}:")
        for g, vals in sorted(groups.items(), key=lambda kv: -len(kv[1]))[:15]:
            lines.append(f"  {g}: {sum(vals) / len(vals):.0%} ({len(vals)})")

    lines.append("\nСредние факторы у «да» и «нет» (фактор без разницы ничего не различает):")
    for fct in FACTORS:
        ys = [float(r[fct]) for r in labeled if r["_label"] and r.get(fct) is not None]
        ns = [float(r[fct]) for r in labeled if not r["_label"] and r.get(fct) is not None]
        if ys and ns:
            lines.append(f"  {fct:<15} да {sum(ys) / len(ys):.2f}   нет {sum(ns) / len(ns):.2f}")

    fp = [r for r in labeled if not r["_label"]][:10]
    if fp:
        lines.append("\nПримеры ошибок (помечены «нет»), начиная с самых высоких:")
        for r in fp:
            lines.append(f"  PQI {r.get('PQI')}: {str(r.get('Цитата') or '')[:110]}")
    return "\n".join(lines)


def main(argv=None) -> int:
    argv = argv if argv is not None else sys.argv[1:]
    if not argv:
        print(__doc__)
        return 1
    print(report(argv[0]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
