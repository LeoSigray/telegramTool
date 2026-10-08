"""Excel-отчёт: люди по убыванию PQI, источники, отсеянные, профиль и лог.

Тексты из чатов — недоверенные данные: строки, начинающиеся с «=», «+», «-», «@»,
записываются как текст, а не как формулы (защита от формул-инъекций в Excel).
"""
from __future__ import annotations

import os
from collections import Counter, defaultdict

from .llm_judge import INTENT_RU, ROLE_RU

TIER_FILL = {"A": "C6EFCE", "B": "FFEB9C", "C": "EDEDED"}
HEADER_FILL = "DDEBF7"
DROP_SAMPLE_PER_REASON = 40

PEOPLE_COLS = [
    ("№", 5), ("Интерес", 8), ("Уровень", 8), ("Тип", 10), ("Username", 20), ("Имя", 22),
    ("User ID", 13),
    # кто этот человек
    ("О себе (био)", 40), ("В сети", 16), ("Premium", 8), ("ЛПР по био", 16),
    # пересечения: где он есть в нише
    ("Чатов ниши: пишет", 10), ("Чатов ниши: состоит", 10), ("Сообщений по теме", 10),
    ("Сообщений за окно", 10), ("Чаты, где пишет", 34), ("Состоит также в", 34),
    ("Другие его запросы", 60),
    # его запрос
    ("Сообщение", 11), ("Цитата", 60), ("Источник", 24), ("Дата (UTC)", 17), ("Возраст, ч", 9),
    ("Тип запроса", 18), ("Проверка", 13),
    # оценка
    ("PQI запроса", 8), ("Близость к нише", 9), ("I намерение", 9), ("R тема", 8), ("F свежесть", 9),
    ("T ветка", 8), ("E активность", 10), ("A ЛПР", 7), ("CQI источника", 10),
    ("Запросов от автора", 10), ("Почему в списке", 70), ("Предупреждения", 40),
    ("Разметка (да/нет)", 12),
]
WRAP_COLS = {"О себе (био)", "Чаты, где пишет", "Состоит также в", "Другие его запросы", "Цитата",
             "Почему в списке", "Предупреждения"}
KIND_RU = {"hot": "горячий", "warm": "тёплый"}
STATUS_RU = {"online": "в сети", "recently": "недавно", "last_week": "на этой неделе",
             "last_month": "в этом месяце", "empty": "скрыт", "": "—"}


def _seen(au) -> str:
    if au is None:
        return "—"
    if au.status == "offline" and au.was_online:
        return f"был {au.was_online[:10]}"
    return STATUS_RU.get(au.status, au.status or "—")


SOURCE_COLS = [
    ("Источник", 26), ("Ссылка", 30), ("Тип", 9), ("Участников", 11), ("Найден через", 34),
    ("Статус", 10), ("Причина", 36), ("Предфильтр", 10), ("Сообщений людей", 11),
    ("Запросов/нед", 11), ("Авторов запросов/нед", 12), ("Доля с username", 11),
    ("Средняя тема", 10), ("Конкуренция", 11), ("Доля рекламы", 10), ("CQI", 9), ("CQI норм.", 9),
    ("Признаки мусора", 60), ("Язык (доля)", 12), ("Активность, сообщ./день", 12),
    ("Проба: запросов по теме/нед", 14), ("Вес чата для интереса", 10),
    ("Доля людей, которые есть в других чатах ниши", 14), ("Участников в списке", 11),
    ("Доля сообщений о товаре клиента", 12),
]

DROP_COLS = [("Причина", 40), ("Автор", 22), ("Источник", 24), ("Ссылка", 11), ("Цитата", 70),
             ("I (правила)", 10), ("R тема", 8)]


def _safe(value):
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub(" ", value)[:32000]
    return value


def _set(ws, row: int, col: int, value, link: str | None = None):
    cell = ws.cell(row=row, column=col, value=_safe(value))
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        cell.data_type = "s"
    if link:
        cell.hyperlink = link
        cell.style = "Hyperlink"
    return cell


def _header(ws, cols) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    for j, (name, width) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=j, value=name)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[cell.column_letter].width = width
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32


def _quote(text: str, n: int = 300) -> str:
    t = " ".join((text or "").split())
    return t if len(t) <= n else t[:n - 1] + "…"


def _people_sheet(ws, people: list) -> None:
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, PatternFill

    _header(ws, PEOPLE_COLS)
    col = {name: j for j, (name, _) in enumerate(PEOPLE_COLS, 1)}
    wrap = Alignment(wrap_text=True, vertical="top")
    for n, c in enumerate(people, 1):
        r = n + 1
        au = c.author
        uname = f"@{au.username}" if au and au.username else ""
        v = c.verdict
        wrote = bool(c.msg.msg_id)
        values = {
            "№": n, "Интерес": round(c.interest, 1), "Уровень": c.tier,
            "Тип": KIND_RU.get(c.kind, c.kind), "Username": uname,
            "Имя": au.name if au else "", "User ID": c.msg.sender_id,
            "Чатов ниши: пишет": c.chats_active, "Чатов ниши: состоит": c.chats_member,
            "Сообщений по теме": c.topic_msgs, "Состоит также в": ", ".join(c.member_chats),
            "PQI запроса": round(c.pqi, 1) if c.kind == "hot" else "",
            "Близость к нише": round(c.affinity, 2),
            "О себе (био)": (au.about or "") if au else "",
            "В сети": _seen(au), "Premium": "да" if au and au.premium else "",
            "ЛПР по био": c.a_hit or ("по оценке LLM" if c.a >= 0.8 else ""),
            "Сообщений за окно": c.n_msgs, "Чаты, где пишет": ", ".join(c.chats),
            "Другие его запросы": "\n".join(c.other_requests),
            "Сообщение": "открыть" if wrote else "",
            "Цитата": _quote(c.msg.text) if wrote else "— в окне не писал, состоит в чатах ниши",
            "Источник": c.source.label(),
            "Дата (UTC)": c.msg.date.strftime("%Y-%m-%d %H:%M") if wrote else "",
            "Возраст, ч": round(c.age_h, 1) if wrote else "",
            "Тип запроса": INTENT_RU.get(c.intent_type, "—"),
            "Проверка": f"LLM: {ROLE_RU.get(v.role, v.role)}" if v else "правила",
            "I намерение": round(c.i, 2) if c.kind == "hot" else "",
            "R тема": round(c.rel, 2), "F свежесть": round(c.fresh, 2),
            "T ветка": round(c.t, 2), "E активность": round(c.e, 2), "A ЛПР": round(c.a, 2),
            "CQI источника": round(c.cqi_n, 2), "Запросов от автора": c.n_requests,
            "Почему в списке": c.why, "Предупреждения": "; ".join(c.warnings), "Разметка (да/нет)": "",
        }
        links = {"Username": f"https://t.me/{au.username}" if uname else None,
                 "Сообщение": c.msg.link(c.source) if wrote else None, "Источник": c.source.link()}
        for name, j in col.items():
            cell = _set(ws, r, j, values[name], links.get(name))
            if name in WRAP_COLS:
                cell.alignment = wrap
        ws.cell(row=r, column=col["Уровень"]).fill = PatternFill("solid", fgColor=TIER_FILL.get(c.tier, "FFFFFF"))
        if c.kind == "hot":
            note = (f"Интерес = 100 × (вес запроса × горячесть {c.hot:.2f} + вес ниши × близость "
                    f"{c.affinity:.2f} × (0.5+0.5·ЛПР {c.a:.2f})) = {c.interest:.1f}\n"
                    f"PQI запроса = 100 × I {c.i:.2f} × R {c.rel:.2f} × F {c.fresh:.2f} × T {c.t:.2f}"
                    f" × E {c.e:.2f} × (0.5+0.5·A {c.a:.2f}) × (0.7+0.3·CQI {c.cqi_n:.2f})"
                    f" × бонус {c.bonus:.2f} = {c.pqi:.1f}")
        else:
            note = (f"Интерес = 100 × вес ниши × близость {c.affinity:.2f} × активность {c.e:.2f}"
                    f" × (0.5+0.5·ЛПР {c.a:.2f}) = {c.interest:.1f}\n"
                    f"Близость: пишет в {c.chats_active} чатах ниши, состоит ещё в {c.chats_member},"
                    f" сообщений по теме {c.topic_msgs}")
        ws.cell(row=r, column=col["Интерес"]).comment = Comment(note, "smart-parser")
    if people:
        ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(PEOPLE_COLS)).column_letter}{len(people) + 1}"


def _sources_sheet(ws, report: list, cqi: dict, junk_metrics: dict | None = None,
                   chat_info: dict | None = None) -> None:
    from .junk import summary
    _header(ws, SOURCE_COLS)
    order = {"selected": 0, "queued": 1, "error": 2, "rejected": 3}
    rows = sorted(report, key=lambda s: (order.get(s.status, 3),
                                          -cqi.get(s.chat_id, {}).get("cqi", 0.0), -s.meta_score))
    status_ru = {"selected": "взят", "queued": "в очереди", "rejected": "отброшен", "error": "ошибка"}
    for r, s in enumerate(rows, 2):
        m = cqi.get(s.chat_id, {})
        vals = [s.title or s.label(), s.link(), "канал" if s.kind == "channel" else "группа",
                s.members, s.found_via, status_ru.get(s.status, s.status), s.reason,
                round(s.meta_score, 3) if s.meta_score < 10 else "вручную",
                m.get("read", ""), _r(m.get("requests_week")), _r(m.get("authors_week")),
                _r(m.get("username_share")), _r(m.get("rel")), _r(m.get("comp")), _r(m.get("ad")),
                _r(m.get("cqi"), 3), _r(m.get("cqi_n")),
                summary((junk_metrics or {}).get(s.chat_id)),
                f"{s.lang} {s.lang_share:.0%}" if s.lang else "", _r(s.activity or None, 1),
                _r(s.y_est if s.lang else None, 1),
                _r((chat_info or {}).get(s.chat_id, {}).get("weight")),
                _r((chat_info or {}).get(s.chat_id, {}).get("overlap")),
                (chat_info or {}).get(s.chat_id, {}).get("members", ""),
                _r((chat_info or {}).get(s.chat_id, {}).get("product_share"), 3)]
        for j, val in enumerate(vals, 1):
            _set(ws, r, j, val, s.link() if j == 2 else None)


def _r(x, nd: int = 2):
    return "" if x is None else round(x, nd)


def _pairs_sheet(ws, pairs: list) -> None:
    _header(ws, [("Чат A", 30), ("Чат B", 30), ("Общих людей", 12), ("Доля от объединения", 12)])
    for r, (a, b, shared, jac) in enumerate(pairs, 2):
        for j, val in enumerate((a, b, shared, round(jac, 3)), 1):
            _set(ws, r, j, val)


def _dropped_sheet(ws, dropped: list) -> None:
    _header(ws, DROP_COLS)
    by_reason: dict = defaultdict(list)
    for c in dropped:
        by_reason[c.drop].append(c)
    counts = Counter({k: len(v) for k, v in by_reason.items()})
    r = 2
    for reason, _ in counts.most_common():
        sample = sorted(by_reason[reason], key=lambda c: -(c.f.intent_rule * max(c.r, 0.01)))
        for c in sample[:DROP_SAMPLE_PER_REASON]:
            au = c.author
            who = f"@{au.username}" if au and au.username else (au.name if au else str(c.msg.sender_id))
            vals = [f"{reason} (всего {counts[reason]})", who, c.source.label(), "открыть",
                    _quote(c.msg.text, 200), round(c.f.intent_rule, 2), round(c.r, 2)]
            for j, val in enumerate(vals, 1):
                _set(ws, r, j, val, c.msg.link(c.source) if j == 4 else None)
            r += 1


def _profile_sheet(ws, profile, params) -> None:
    from openpyxl.styles import Alignment, Font

    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 110
    rows = [
        ("ПРОФИЛЬ КЛИЕНТА", ""),
        ("Канал", f"@{profile.channel}"),
        ("Название", profile.title),
        ("Описание", profile.about),
        ("Бриф", profile.brief),
        ("Оффер", profile.offer),
        ("Аудитория", profile.audience),
        ("Ключевые термины", ", ".join(profile.topic_terms)),
        ("Фразы покупателя", "\n".join(profile.buyer_phrases)),
        ("Поисковые запросы", ", ".join(profile.search_queries)),
        ("Анти-портрет", "\n".join(profile.anti)),
        ("Фразы продавцов", "\n".join(profile.seller_phrases)),
        ("Построен", profile.built_with),
        ("Хеш профиля", profile.hash()),
        ("", ""),
        ("ПАРАМЕТРЫ ЗАПУСКА", ""),
    ] + [(k, str(v)) for k, v in params.to_dict().items()]
    for i, (k, v) in enumerate(rows, 1):
        a = _set(ws, i, 1, k)
        b = _set(ws, i, 2, v)
        b.alignment = Alignment(wrap_text=True, vertical="top")
        if k.isupper():
            a.font = Font(bold=True)


def _log_sheet(ws, log) -> None:
    from openpyxl.styles import Font

    for col, w in zip("ABCD", (20, 8, 16, 120)):
        ws.column_dimensions[col].width = w
    ws.cell(row=1, column=1, value="ВОРОНКА").font = Font(bold=True)
    r = 2
    for name, val in log.funnel.items():
        _set(ws, r, 1, name)
        _set(ws, r, 3, val)
        r += 1
    r += 1
    for j, name in enumerate(("Время", "Уровень", "Этап", "Сообщение"), 1):
        ws.cell(row=r, column=j, value=name).font = Font(bold=True)
    for row in log.rows:
        r += 1
        for j, val in enumerate(row, 1):
            _set(ws, r, j, val)


def write_report(path: str, people: list, no_username: list, dropped: list, sources_report: list,
                 cqi: dict, profile, params, log, junk_metrics: dict | None = None,
                 reserve: list | None = None, chat_info: dict | None = None,
                 pairs: list | None = None) -> str:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Люди"
    _people_sheet(ws, people)
    if reserve:
        _people_sheet(wb.create_sheet("Запас"), reserve)
    _people_sheet(wb.create_sheet("Без username"), no_username)
    _sources_sheet(wb.create_sheet("Источники"), sources_report, cqi, junk_metrics, chat_info)
    _pairs_sheet(wb.create_sheet("Пересечения чатов"), pairs or [])
    _dropped_sheet(wb.create_sheet("Отсеянные"), dropped)
    _profile_sheet(wb.create_sheet("Профиль"), profile, params)
    _log_sheet(wb.create_sheet("Лог"), log)

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:
        # файл открыт в Excel — сохраняем рядом
        root, ext = os.path.splitext(path)
        path = f"{root}_1{ext}"
        wb.save(path)
    return path
